#!/usr/bin/env python3
"""Flatten a CMS cross-section DB (JSON) into an XLSX (with live sigma_eff formula)
and a LaTeX-rendered PDF. Reusable: `python xsec_view.py db.json [outdir]`.

XLSX columns kept "live": sigma_eff = cross_section_fb * br * kfactor is an Excel
formula, so changing kfactor (e.g. once AN k-factors are confirmed) auto-updates.
"""
from __future__ import annotations

import json
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# --- column schema: (json_key, header, width, kind) ------------------------------
COLUMNS: list[tuple[str, str, int, str]] = [
    ("sample", "sample", 26, "text"),
    ("group", "group", 13, "text"),
    ("accuracy", "accuracy", 12, "text"),
    ("cross_section_fb", "sigma_fb", 15, "sci"),
    ("br", "BR", 11, "num"),
    ("kfactor", "kfactor", 9, "num"),
    ("sigma_eff", "sigma_eff_fb", 15, "formula"),  # = sigma * BR * k
    ("kfactor_an_ref", "k_AN_ref", 9, "num"),
    ("nevents", "n_events", 12, "int"),
    ("nfiles", "n_files", 8, "int"),
    ("frac_neg_weight", "f_negW", 9, "num"),
    ("verify_xsec", "verify", 7, "bool"),
    ("xsec_ref", "xsec_ref", 20, "text"),
    ("das_path", "DAS", 60, "text"),
]


@dataclass(frozen=True)
class Theme:
    header_bg = "1F3864"
    header_fg = "FFFFFF"
    data_bg = "FFF2CC"   # verify_xsec=true highlight
    sig_bg = "E2EFDA"    # signal rows
    font = "Arial"


def classify(name: str, rec: dict) -> str:
    """Assign a process group for filtering/grouping."""
    if rec.get("cross_section_fb") is None and rec.get("br") is None:
        return "DATA"
    n = name.lower()
    if name in {"ttHH"}:
        return "SIGNAL"
    if n.startswith("qcd"):
        return "QCD"
    if name.startswith(("TTTo", "tt4b", "ttbb")):
        return "TTbar"
    if name.startswith(("WJets", "ZJets", "DYJets")):
        return "V+jets"
    if name.startswith(("ST_", "tHq", "tHW")):
        return "SingleTop"
    if name in {"WW", "WZ", "ZZ"}:
        return "Diboson"
    if name.startswith("tt") or name.startswith("ttH"):
        return "ttX"
    return "other"


def load(db_path: Path) -> tuple[dict, list[dict]]:
    raw = json.loads(db_path.read_text())
    meta = raw.pop("_meta", {})
    rows = []
    for name, rec in raw.items():
        rec = dict(rec)
        rec["sample"] = name
        rec["group"] = classify(name, rec)
        rows.append(rec)
    # MC first (by group then descending xsec), data last
    rows.sort(key=lambda r: (r["group"] == "DATA", r["group"],
                             -(r.get("cross_section_fb") or 0)))
    return meta, rows


# --- XLSX -----------------------------------------------------------------------
def write_xlsx(meta: dict, rows: list[dict], out: Path) -> None:
    th = Theme()
    wb = Workbook()
    ws = wb.active
    ws.title = "xsec_db"

    col_of = {key: i + 1 for i, (key, *_rest) in enumerate(COLUMNS)}
    xs_c, br_c, kf_c = (get_column_letter(col_of[k])
                        for k in ("cross_section_fb", "br", "kfactor"))

    # header
    for c, (_key, hdr, width, _kind) in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, c, hdr)
        cell.font = Font(name=th.font, bold=True, color=th.header_fg, size=10)
        cell.fill = PatternFill("solid", fgColor=th.header_bg)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width

    for r, rec in enumerate(rows, start=2):
        is_data = rec["group"] == "DATA"
        for c, (key, _hdr, _w, kind) in enumerate(COLUMNS, start=1):
            if key == "sigma_eff":
                val = (None if is_data
                       else f"={xs_c}{r}*{br_c}{r}*{kf_c}{r}")
            else:
                val = rec.get(key)
            cell = ws.cell(r, c, val)
            cell.font = Font(name=th.font, size=9)
            if kind == "sci" and isinstance(val, (int, float)):
                cell.number_format = "0.000E+00"
            elif kind == "formula":
                cell.number_format = "0.000E+00"
            elif kind == "num" and isinstance(val, (int, float)):
                cell.number_format = "0.0000"
            elif kind == "int" and isinstance(val, (int, float)):
                cell.number_format = "#,##0"
            elif kind == "bool":
                cell.value = "" if val is None else ("Y" if val else "n")
                cell.alignment = Alignment(horizontal="center")
            if key in ("das_path", "xsec_ref"):
                cell.alignment = Alignment(horizontal="left")
        # row tinting
        fill = None
        if rec["group"] == "SIGNAL":
            fill = PatternFill("solid", fgColor=th.sig_bg)
        elif rec.get("verify_xsec"):
            fill = PatternFill("solid", fgColor=th.data_bg)
        if fill:
            for c in range(1, len(COLUMNS) + 1):
                ws.cell(r, c).fill = fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    # meta sheet
    mws = wb.create_sheet("_meta")
    mws.column_dimensions["A"].width = 18
    mws.column_dimensions["B"].width = 110
    for i, (k, v) in enumerate(meta.items(), start=1):
        mws.cell(i, 1, k).font = Font(name=th.font, bold=True, size=9)
        mws.cell(i, 2, str(v)).font = Font(name=th.font, size=9)
        mws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out)


# --- LaTeX/PDF ------------------------------------------------------------------
def _tex_escape(s: str) -> str:
    return (str(s).replace("\\", r"\textbackslash{}").replace("_", r"\_")
            .replace("%", r"\%").replace("&", r"\&").replace("#", r"\#")
            .replace("^", r"\^{}").replace("~", r"\~{}"))


def _fmt(v, kind: str) -> str:
    if v is None:
        return "--"
    if kind == "sci":
        return f"{v:.3g}"
    if kind == "num":
        return f"{v:.4g}"
    if kind == "int":
        return f"{int(v):,}"
    return _tex_escape(v)


def write_pdf(meta: dict, rows: list[dict], outdir: Path) -> Path | None:
    # compact PDF table: a useful physics subset of columns
    cols = [("sample", "sample", "text"), ("group", "grp", "text"),
            ("accuracy", "acc.", "text"), ("cross_section_fb", r"$\sigma$ [fb]", "sci"),
            ("br", "BR", "num"), ("kfactor", "k", "num"),
            ("sigma_eff", r"$\sigma_{\mathrm{eff}}$ [fb]", "sci"),
            ("nevents", "$N$", "int"), ("frac_neg_weight", "$f_{-}$", "num"),
            ("verify_xsec", "vf", "text")]
    header = " & ".join(h for _k, h, _t in cols) + r" \\"
    body = []
    for rec in rows:
        cells = []
        for key, _h, kind in cols:
            if key == "sigma_eff":
                xs, br, kf = (rec.get("cross_section_fb"), rec.get("br"),
                              rec.get("kfactor"))
                v = (xs * br * kf if None not in (xs, br, kf) else None)
                cells.append(_fmt(v, "sci"))
            elif key == "verify_xsec":
                cells.append("Y" if rec.get("verify_xsec") else "")
            else:
                cells.append(_fmt(rec.get(key), kind))
        line = " & ".join(cells) + r" \\"
        if rec["group"] == "SIGNAL":
            line = r"\rowcolor{green!12}" + line
        elif rec.get("verify_xsec"):
            line = r"\rowcolor{yellow!18}" + line
        elif rec["group"] == "DATA":
            line = r"\rowcolor{gray!10}" + line
        body.append(line)

    tex = r"""\documentclass[9pt]{extarticle}
\usepackage[a4paper,landscape,margin=1.2cm]{geometry}
\usepackage{booktabs,longtable}
\usepackage[table]{xcolor}
\usepackage[T1]{fontenc}
\renewcommand{\arraystretch}{1.18}
\setlength{\tabcolsep}{4pt}
\begin{document}
\begin{center}{\large\bfseries CMS cross-section DB --- %(era)s (%(campaign)s)}\\[2pt]
{\footnotesize $\mathcal{L}=%(lumi).2f$\,fb$^{-1}$ \quad
$\sigma_{\mathrm{eff}}=\sigma\times\mathrm{BR}\times k$ \quad
\colorbox{green!12}{signal}\;\colorbox{yellow!18}{verify\_xsec}\;\colorbox{gray!10}{data}}\end{center}
\vspace{2pt}
\footnotesize
\begin{longtable}{lll r r r r r r c}
\toprule
%(header)s
\midrule
\endhead
\bottomrule
\endfoot
%(body)s
\end{longtable}
\end{document}
""" % {
        "era": _tex_escape(meta.get("era", "")),
        "campaign": _tex_escape(meta.get("campaign", "")),
        "lumi": float(meta.get("lumi_fb_inv", 0) or 0),
        "header": header,
        "body": "\n".join(body),
    }
    texf = outdir / "xsec_db.tex"
    texf.write_text(tex)
    for _ in range(2):  # longtable needs 2 passes
        proc = subprocess.run(
            ["pdflatex", "-interaction=nonstopmode", "-halt-on-error", texf.name],
            cwd=outdir, capture_output=True, text=True)
    pdf = outdir / "xsec_db.pdf"
    if not pdf.exists():
        sys.stderr.write(proc.stdout[-2000:])
        return None
    return pdf


def main() -> None:
    db = Path(sys.argv[1])
    outdir = Path(sys.argv[2]) if len(sys.argv) > 2 else db.parent
    outdir.mkdir(parents=True, exist_ok=True)
    meta, rows = load(db)
    xlsx = outdir / "xsec_db.xlsx"
    write_xlsx(meta, rows, xlsx)
    print(f"XLSX -> {xlsx}  ({len(rows)} samples)")
    pdf = write_pdf(meta, rows, outdir)
    print(f"PDF  -> {pdf}" if pdf else "PDF  -> FAILED (see stderr)")


if __name__ == "__main__":
    main()
